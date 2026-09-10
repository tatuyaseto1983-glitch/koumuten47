#!/usr/bin/env python3
"""社人研「日本の地域別将来推計人口（令和5年推計）」を取り込む。

e-Stat APIには収載がないため、社人研サイトの結果表（Excel）を直接読み込み、
estat_fetch.py と同じ列構成のCSVに変換する。2020年を起点に2050年までの
5年ごとの推計値が、都道府県と市区町村の両方について入っている。

  python3 scripts/ipss_fetch.py            # ダウンロードして変換
  python3 scripts/ipss_fetch.py --no-download   # 手元のファイルだけで変換
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
import time
from pathlib import Path

import openpyxl
import requests

BASE_URL = "https://www.ipss.go.jp/pp-shicyoson/j/shicyoson23/2gaiyo_hyo"
ROOT = Path(__file__).resolve().parent.parent
XLSX_DIR = ROOT / "data" / "ipss"
OUT_DIR = ROOT / "data" / "processed"

# 実数の表のみを取り込む（割合は人数から算出できるため、まずは実数を揃える）
FILES = [
    "kekkahyo1.xlsx",    # 総人口
    "kekkahyo2_1.xlsx",  # 0〜14歳人口
    "kekkahyo2_2.xlsx",  # 15〜64歳人口
    "kekkahyo2_3.xlsx",  # 65歳以上人口
    "kekkahyo2_4.xlsx",  # 75歳以上人口
]

FIELDS = ["indicator", "indicator_label", "area_code", "area_name", "area_level",
          "pref_code", "pref_name", "time_code", "year", "cat_code", "cat_name",
          "unit", "value", "value_raw"]


def download(files: list[str], sleep: float = 2.0) -> None:
    XLSX_DIR.mkdir(parents=True, exist_ok=True)
    for name in files:
        dest = XLSX_DIR / name
        if dest.exists():
            print(f"  {name} は既にあるので飛ばします。")
            continue
        print(f"  ダウンロード中: {name}")
        res = requests.get(f"{BASE_URL}/{name}", timeout=120)
        res.raise_for_status()
        dest.write_bytes(res.content)
        time.sleep(sleep)


def table_label(sheet_title: str) -> str:
    """A1の見出しから「総人口」「0～14歳人口」などの指標名を取り出す。"""
    text = str(sheet_title or "").replace("\n", " ")
    # 「結果表2-1　0～14歳人口および指数（…」→「0～14歳人口」
    m = re.search(r"結果表\S+\s+(.+?)および指数", text)
    if m:
        return m.group(1).strip()
    m = re.search(r"結果表\S+\s+(.+)", text)
    return m.group(1).strip()[:20] if m else text.strip()[:20]


def parse_file(path: Path) -> tuple[str, list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    label = table_label(str(rows[0][0]) if rows and rows[0] else "")

    header_idx = next(i for i, r in enumerate(rows)
                      if r and str(r[0]).strip() == "コード")
    # 見出しの次の行に「2020年」「2025年」…が並ぶ
    year_row = rows[header_idx + 1]
    year_cols: list[tuple[int, str]] = []
    for col, cell in enumerate(year_row):
        m = re.match(r"(\d{4})年", str(cell or ""))
        if m and col >= 4:
            year_cols.append((col, m.group(1)))
    # 実数のあとに「指数」の同じ年が続くため、最初のひとまわりだけ使う
    seen: set[str] = set()
    trimmed = []
    for col, year in year_cols:
        if year in seen:
            break
        seen.add(year)
        trimmed.append((col, year))
    year_cols = trimmed

    out: list[dict] = []
    skipped = 0
    for r in rows[header_idx + 2:]:
        if not r or r[0] in (None, ""):
            continue
        code = str(r[0]).strip()
        if not code.isdigit():
            continue
        area_code = code.zfill(5)
        kind = str(r[1] or "").strip()
        pref_name = str(r[2] or "").strip()
        city_name = str(r[3] or "").strip()
        level = "pref" if kind == "a" else "city"
        area_name = pref_name if level == "pref" else city_name
        for col, year in year_cols:
            raw = r[col] if col < len(r) else None
            try:
                value = float(raw)
            except (TypeError, ValueError):
                value = None
                if raw not in (None, ""):
                    skipped += 1
            out.append({
                "indicator": "future_population",
                "indicator_label": "将来推計人口（社人研 令和5年推計）",
                "area_code": area_code,
                "area_name": area_name,
                "area_level": level,
                "pref_code": area_code[:2],
                "pref_name": pref_name,
                "time_code": f"{year}000000",
                "year": year,
                "cat_code": path.stem,
                "cat_name": label,
                "unit": "人",
                "value": "" if value is None else value,
                "value_raw": "" if raw is None else str(raw),
            })
    if skipped:
        print(f"    数値として読めなかったセル: {skipped}件（秘匿・非該当）")
    return label, out


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)


def main() -> None:
    ap = argparse.ArgumentParser(description="社人研 将来推計人口の取り込み")
    ap.add_argument("--no-download", action="store_true",
                    help="ダウンロードせず data/ipss/ の既存ファイルだけを使う")
    args = ap.parse_args()

    if not args.no_download:
        print("社人研サイトから結果表を取得します。")
        download(FILES)

    all_rows: list[dict] = []
    for name in FILES:
        path = XLSX_DIR / name
        if not path.exists():
            print(f"  {name} がありません。飛ばします。", file=sys.stderr)
            continue
        label, rows = parse_file(path)
        print(f"  {name}: {label}  {len(rows):,}行")
        all_rows.extend(rows)

    if not all_rows:
        raise SystemExit("取り込めるデータがありませんでした。")

    for level, suffix in (("pref", "pref"), ("city", "city")):
        subset = [r for r in all_rows if r["area_level"] == level]
        out = OUT_DIR / f"future_population_{suffix}.csv"
        write_csv(out, subset)
        areas = len({r["area_code"] for r in subset})
        years = sorted({r["year"] for r in subset})
        print(f"→ {out.relative_to(ROOT)}  {len(subset):,}行 / 地域{areas}件 / "
              f"{years[0]}〜{years[-1]}")


if __name__ == "__main__":
    main()
